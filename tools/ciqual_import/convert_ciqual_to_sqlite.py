#!/usr/bin/env python3
"""Convert the CIQUAL nutrition table (xlsx) into a lean SQLite database
optimized for the Flutter app's offline food search + meal logging feature.

Usage:
    python3 -m venv .venv && source .venv/bin/activate
    pip install -r requirements.txt
    python3 convert_ciqual_to_sqlite.py

Reads:  ../../food_bd/ciqual/Table Ciqual 2025_ENG_2025_11_03.xlsx (sheet1)
Writes: ../../assets/db/nutrition.db
"""
from __future__ import annotations

import re
import sqlite3
import sys
from pathlib import Path

from openpyxl import load_workbook

from text_normalizer import normalize

ROOT = Path(__file__).resolve().parents[2]
SOURCE_XLSX = ROOT / "food_bd" / "ciqual" / "Table Ciqual 2025_ENG_2025_11_03.xlsx"
OUTPUT_DB = ROOT / "assets" / "db" / "nutrition.db"

# Header text is used verbatim from the sheet (with embedded newlines), so we
# match on normalized (lowercased, whitespace-collapsed) substrings instead of
# exact strings — this keeps the script robust to minor header rewording in
# future CIQUAL releases.
COLUMN_MATCHERS = {
    "id": lambda h: h == "alim_code",
    "food_name": lambda h: h == "alim_nom_eng",
    "calories_kcal_100g": lambda h: "kcal" in h and "1169" in h,
    "protein_g_100g": lambda h: h.startswith("protein") and "crude" not in h,
    "carbs_g_100g": lambda h: h.startswith("carbohydrate"),
    "fat_g_100g": lambda h: h.startswith("fat "),
    "fiber_g_100g": lambda h: h.startswith("fibres"),
}


def normalize_header(raw: str | None) -> str:
    if raw is None:
        return ""
    return re.sub(r"\s+", " ", str(raw).strip().lower())


def find_column_indices(header_row: list) -> dict[str, int]:
    normalized = [normalize_header(cell) for cell in header_row]
    indices: dict[str, int] = {}
    for field, matcher in COLUMN_MATCHERS.items():
        matches = [i for i, h in enumerate(normalized) if matcher(h)]
        if len(matches) != 1:
            raise ValueError(
                f"Expected exactly 1 header match for '{field}', got {len(matches)}: "
                f"{[header_row[i] for i in matches]}"
            )
        indices[field] = matches[0]
    return indices


def parse_float(value) -> float | None:
    if value is None:
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip()
    if not text or text in {"-", "traces", "traces polyols"}:
        return None
    text = text.replace(",", ".")
    try:
        return float(text)
    except ValueError:
        return None


def normalize_food_name(raw: str) -> str:
    return re.sub(r"\s+", " ", str(raw).strip())


def read_foods(xlsx_path: Path) -> list[dict]:
    wb = load_workbook(xlsx_path, read_only=True, data_only=True)
    sheet = wb["Sheet1"] if "Sheet1" in wb.sheetnames else wb.worksheets[0]
    rows = sheet.iter_rows(values_only=True)
    header_row = next(rows)
    indices = find_column_indices(list(header_row))

    foods = []
    for row in rows:
        if row is None or all(c is None for c in row):
            continue
        raw_id = row[indices["id"]]
        raw_name = row[indices["food_name"]]
        if raw_id is None or raw_name is None:
            continue
        foods.append(
            {
                "id": int(raw_id) if isinstance(raw_id, (int, float)) else int(str(raw_id).strip()),
                "food_name": normalize_food_name(raw_name),
                "calories_kcal_100g": parse_float(row[indices["calories_kcal_100g"]]),
                "protein_g_100g": parse_float(row[indices["protein_g_100g"]]),
                "carbs_g_100g": parse_float(row[indices["carbs_g_100g"]]),
                "fat_g_100g": parse_float(row[indices["fat_g_100g"]]),
                "fiber_g_100g": parse_float(row[indices["fiber_g_100g"]]),
            }
        )
    wb.close()
    return foods


def dedupe(foods: list[dict]) -> list[dict]:
    seen: dict[str, dict] = {}
    dropped = 0
    for food in foods:
        key = normalize(food["food_name"])
        if key in seen:
            dropped += 1
            continue
        seen[key] = food
    if dropped:
        print(f"Dropped {dropped} duplicate food name(s) after normalization.")
    return list(seen.values())


SCHEMA = """
CREATE TABLE foods (
    id INTEGER PRIMARY KEY,
    food_name TEXT NOT NULL,
    search_name TEXT NOT NULL,
    calories_kcal_100g REAL,
    protein_g_100g REAL,
    carbs_g_100g REAL,
    fat_g_100g REAL,
    fiber_g_100g REAL
);
CREATE INDEX idx_foods_search_name ON foods(search_name);

CREATE TABLE meal_entries (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    food_id INTEGER NOT NULL REFERENCES foods(id),
    food_name TEXT NOT NULL,
    grams REAL NOT NULL,
    calories REAL NOT NULL,
    protein REAL NOT NULL,
    carbs REAL NOT NULL,
    fat REAL NOT NULL,
    logged_at TEXT NOT NULL,
    log_date TEXT NOT NULL
);
CREATE INDEX idx_meal_entries_log_date ON meal_entries(log_date);
"""


def write_database(foods: list[dict], output_path: Path) -> None:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    if output_path.exists():
        output_path.unlink()

    conn = sqlite3.connect(output_path)
    try:
        conn.executescript(SCHEMA)
        conn.executemany(
            """
            INSERT INTO foods (
                id, food_name, search_name, calories_kcal_100g,
                protein_g_100g, carbs_g_100g, fat_g_100g, fiber_g_100g
            ) VALUES (:id, :food_name, :search_name, :calories_kcal_100g,
                      :protein_g_100g, :carbs_g_100g, :fat_g_100g, :fiber_g_100g)
            """,
            [{**food, "search_name": normalize(food["food_name"])} for food in foods],
        )
        conn.commit()
        conn.execute("VACUUM")
    finally:
        conn.close()


def main() -> None:
    if not SOURCE_XLSX.exists():
        print(f"Source file not found: {SOURCE_XLSX}", file=sys.stderr)
        sys.exit(1)

    print(f"Reading {SOURCE_XLSX} ...")
    foods = read_foods(SOURCE_XLSX)
    print(f"Parsed {len(foods)} rows.")

    foods = dedupe(foods)
    print(f"{len(foods)} unique foods after de-duplication.")

    write_database(foods, OUTPUT_DB)
    print(f"Wrote {OUTPUT_DB} ({OUTPUT_DB.stat().st_size / 1024:.0f} KB).")


if __name__ == "__main__":
    main()
